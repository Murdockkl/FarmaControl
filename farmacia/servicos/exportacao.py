"""Exporta um Relatorio para CSV, Excel (.xlsx) ou PDF."""
from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape

from ..dominio import FarmaciaError
from ..texto import fmt_data, fmt_data_hora, fmt_numero_decimal
from .relatorios import Relatorio, formatar_valor

_TIPOS_NUMERICOS = {"inteiro", "sinal", "dinheiro"}


# -------------------------------------------------------------------------- CSV
def exportar_csv(rel: Relatorio, caminho: str | Path, nome_farmacia: str = "") -> Path:
    """CSV com ';' e UTF-8 com BOM: abre direto no Excel em português."""
    caminho = Path(caminho)
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([c.titulo + (" (R$)" if c.tipo == "dinheiro" else "") for c in rel.colunas])
        for linha in rel.linhas:
            w.writerow([_valor_csv(v, c.tipo) for v, c in zip(linha, rel.colunas)])
    return caminho


def _valor_csv(valor, tipo: str) -> str:
    if valor is None:
        return ""
    if tipo == "dinheiro":
        return fmt_numero_decimal(valor)
    if tipo in ("inteiro", "sinal"):
        return str(valor)
    return formatar_valor(valor, tipo)


# ------------------------------------------------------------------------ Excel
def exportar_excel(rel: Relatorio, caminho: str | Path, nome_farmacia: str = "") -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise FarmaciaError("Para exportar para Excel, instale a biblioteca openpyxl:\n"
                            "pip install openpyxl") from None

    caminho = Path(caminho)
    wb = Workbook()
    ws = wb.active
    ws.title = rel.titulo[:31]

    ws["A1"] = nome_farmacia or "FarmaControl"
    ws["A1"].font = Font(bold=True, size=11, color="0F766E")
    ws["A2"] = rel.titulo
    ws["A2"].font = Font(bold=True, size=15)
    ws["A3"] = rel.subtitulo
    ws["A3"].font = Font(italic=True, color="666666")

    linha_cab = 5
    preenchimento = PatternFill("solid", fgColor="0F766E")
    borda = Border(bottom=Side(style="thin", color="D5DEDC"))
    for i, col in enumerate(rel.colunas, start=1):
        cel = ws.cell(row=linha_cab, column=i, value=col.titulo)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = preenchimento
        cel.alignment = Alignment(horizontal="right" if col.tipo in _TIPOS_NUMERICOS else "left",
                                  vertical="center", wrap_text=True)

    larguras = [len(c.titulo) + 2 for c in rel.colunas]
    for r, linha in enumerate(rel.linhas, start=linha_cab + 1):
        for i, (valor, col) in enumerate(zip(linha, rel.colunas), start=1):
            if col.tipo == "dinheiro" and valor is not None:
                valor = valor / 100
            cel = ws.cell(row=r, column=i, value=valor)
            cel.border = borda
            if col.tipo == "dinheiro":
                cel.number_format = '"R$" #,##0.00'
            elif col.tipo == "data":
                cel.number_format = "DD/MM/YYYY"
            elif col.tipo == "datahora":
                cel.number_format = "DD/MM/YYYY HH:MM"
            elif col.tipo == "sinal":
                cel.number_format = "+#,##0;-#,##0;0"
            elif col.tipo == "inteiro":
                cel.number_format = "#,##0"
            texto = formatar_valor(linha[i - 1], col.tipo)
            larguras[i - 1] = max(larguras[i - 1], min(len(texto) + 2, 55))

    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[linha_cab].height = 30
    ws.freeze_panes = ws.cell(row=linha_cab + 1, column=1)
    if rel.linhas:
        ws.auto_filter.ref = (f"A{linha_cab}:{get_column_letter(len(rel.colunas))}"
                              f"{linha_cab + len(rel.linhas)}")

    r = linha_cab + len(rel.linhas) + 2
    for rotulo, valor in rel.resumo:
        ws.cell(row=r, column=1, value=rotulo).font = Font(bold=True)
        ws.cell(row=r, column=2, value=valor)
        r += 1
    wb.save(caminho)
    return caminho


# -------------------------------------------------------------------------- PDF
def _pdf_seguro(texto: str) -> str:
    """As fontes padrão do PDF só têm caracteres Latin-1: troca o que não existir nelas."""
    return str(texto).encode("cp1252", "replace").decode("cp1252")


def exportar_pdf(rel: Relatorio, caminho: str | Path, nome_farmacia: str = "") -> Path:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        raise FarmaciaError("Para exportar para PDF, instale a biblioteca reportlab:\n"
                            "pip install reportlab") from None

    caminho = Path(caminho)
    verde = colors.HexColor("#0F766E")
    pagina = landscape(A4)
    margem = 12 * mm
    doc = SimpleDocTemplate(str(caminho), pagesize=pagina, leftMargin=margem, rightMargin=margem,
                            topMargin=margem, bottomMargin=14 * mm,
                            title=rel.titulo, author=nome_farmacia or "FarmaControl")
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=estilos["Title"], fontSize=17, alignment=0,
                            textColor=colors.HexColor("#1F2D2B"), spaceAfter=2)
    sub = ParagraphStyle("s", parent=estilos["Normal"], fontSize=9, textColor=colors.grey)
    empresa = ParagraphStyle("e", parent=estilos["Normal"], fontSize=10, textColor=verde,
                             fontName="Helvetica-Bold")
    celula = ParagraphStyle("c", parent=estilos["Normal"], fontSize=7.5, leading=9)
    celula_dir = ParagraphStyle("cd", parent=celula, alignment=2)
    cabecalho = ParagraphStyle("h", parent=celula, textColor=colors.white,
                               fontName="Helvetica-Bold")
    cabecalho_dir = ParagraphStyle("hd", parent=cabecalho, alignment=2)

    def p(texto, estilo):
        return Paragraph(escape(_pdf_seguro(texto)), estilo)

    elementos = [p(nome_farmacia or "FarmaControl", empresa), p(rel.titulo, titulo),
                 p(rel.subtitulo, sub), Spacer(1, 5 * mm)]

    largura_total = pagina[0] - 2 * margem
    soma = sum(c.largura for c in rel.colunas) or 1
    larguras = [largura_total * c.largura / soma for c in rel.colunas]
    dados = [[p(c.titulo, cabecalho_dir if c.tipo in _TIPOS_NUMERICOS else cabecalho)
              for c in rel.colunas]]
    for linha in rel.linhas_formatadas():
        dados.append([p(v, celula_dir if c.tipo in _TIPOS_NUMERICOS else celula)
                      for v, c in zip(linha, rel.colunas)])
    if not rel.linhas:
        dados.append([p("Nenhum registro encontrado.", celula)] + [""] * (len(rel.colunas) - 1))

    tabela = Table(dados, colWidths=larguras, repeatRows=1)
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), verde),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F7F6")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DEDC")),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    elementos.append(tabela)
    if rel.resumo:
        elementos.append(Spacer(1, 6 * mm))
        for rotulo, valor in rel.resumo:
            elementos.append(p(f"{rotulo}: {valor}", ParagraphStyle(
                "r", parent=estilos["Normal"], fontSize=9, fontName="Helvetica-Bold")))

    def rodape(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(margem, 8 * mm, _pdf_seguro(f"{nome_farmacia or 'FarmaControl'} - {rel.titulo}"))
        canvas.drawRightString(pagina[0] - margem, 8 * mm, f"Página {documento.page}")
        canvas.restoreState()

    doc.build(elementos, onFirstPage=rodape, onLaterPages=rodape)
    return caminho
