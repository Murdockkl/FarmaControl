"""Testes da lógica do sistema (sem interface gráfica).

Cobrem os 15 passos de verificação pedidos, além de validações, exportações e demonstração.
Execute com:  python -m pytest -v
"""
from __future__ import annotations

import sqlite3
from datetime import date, datetime, timedelta

import pytest

from farmacia.contexto import Contexto
from farmacia.dominio import (ALERTA_ESTOQUE_BAIXO, ALERTA_PROXIMO_VENCIMENTO,
                              ALERTA_SEM_ESTOQUE, ALERTA_VENCIDO, BackupError, DuplicadoError,
                              EstoqueInsuficienteError, MedicamentoDados, NaoEncontradoError,
                              ValidacaoError)
from farmacia.servicos import exportacao
from farmacia.servicos.medicamentos import SIT_BAIXO, SIT_PROXIMO, SIT_VENCIDO
from farmacia.servicos.relatorios import (REL_ESTOQUE_ATUAL, REL_ESTOQUE_BAIXO,
                                          REL_MOVIMENTACOES, REL_PROXIMOS, REL_VALOR,
                                          REL_VENCIDOS)
from farmacia.validacao import converter_dinheiro, converter_data, converter_inteiro

HOJE = date.today()


@pytest.fixture
def caminho(tmp_path):
    return tmp_path / "dados" / "farmacia.db"


@pytest.fixture
def ctx(caminho):
    c = Contexto(caminho)
    yield c
    c.fechar()


def paracetamol(**extra) -> MedicamentoDados:
    base = dict(nome="Paracetamol 750 mg", principio_ativo="Paracetamol", eh_generico=True,
                nome_generico="Paracetamol 750 mg Genérico", laboratorio="Lab Teste",
                categoria="Analgésico", dosagem="750", unidade_medida="mg",
                forma_farmaceutica="Comprimido", quantidade_atual=20, estoque_minimo=10,
                preco_compra_centavos=420, preco_venda_centavos=890, setor="Medicamentos",
                corredor="A", estante="02", prateleira="03", lote="L001",
                data_fabricacao=date(2025, 1, 10), data_validade=HOJE + timedelta(days=400))
    base.update(extra)
    return MedicamentoDados(**base)


def ibuprofeno(**extra) -> MedicamentoDados:
    base = dict(nome="Ibuprofeno 600 mg", principio_ativo="Ibuprofeno", laboratorio="Lab Teste",
                categoria="Anti-inflamatório", dosagem="600", unidade_medida="mg",
                forma_farmaceutica="Comprimido", quantidade_atual=40, estoque_minimo=15,
                preco_compra_centavos=500, preco_venda_centavos=1100, setor="Medicamentos",
                corredor="B", estante="01", prateleira="01", lote="L100",
                data_validade=HOJE + timedelta(days=300))
    base.update(extra)
    return MedicamentoDados(**base)


# =========================================================== os 15 passos pedidos
def test_passos_1_a_4_cadastro_persiste_apos_fechar_e_reabrir(caminho):
    ctx = Contexto(caminho)
    med_id = ctx.medicamentos.criar(paracetamol())          # 1. cadastrar
    ctx.fechar()                                            # 2. fechar o sistema

    ctx2 = Contexto(caminho)                                # 3. abrir novamente
    try:
        med = ctx2.medicamentos.obter(med_id)               # 4. continua cadastrado?
        assert med is not None
        assert med.nome == "Paracetamol 750 mg"
        assert med.principio_ativo == "Paracetamol"
        assert med.quantidade_atual == 20
        assert med.laboratorio == "Lab Teste" and med.categoria == "Analgésico"
        assert med.preco_compra_centavos == 420 and med.preco_venda_centavos == 890
        assert (med.corredor, med.estante, med.prateleira) == ("A", "02", "03")
        assert med.data_validade == HOJE + timedelta(days=400)
        assert med.eh_generico is True
        assert len(ctx2.medicamentos.listar()) == 1
    finally:
        ctx2.fechar()


def test_passos_5_6_7_entrada_saida_e_saida_maior_que_estoque(ctx):
    med_id = ctx.medicamentos.criar(paracetamol(quantidade_atual=20))

    entrada = ctx.estoque.entrada(med_id, 30, "Compra")                      # 5. adicionar
    assert (entrada.estoque_anterior, entrada.estoque_posterior) == (20, 50)
    assert ctx.medicamentos.obter(med_id).quantidade_atual == 50

    saida = ctx.estoque.saida(med_id, 5, "Venda")                            # 6. retirar
    assert (saida.estoque_anterior, saida.estoque_posterior) == (50, 45)
    assert ctx.medicamentos.obter(med_id).quantidade_atual == 45

    antes = len(ctx.estoque.listar())
    with pytest.raises(EstoqueInsuficienteError) as erro:                    # 7. retirar 60
        ctx.estoque.saida(med_id, 60, "Venda")
    assert "60 unidades" in str(erro.value) and "45 unidades" in str(erro.value)
    assert ctx.medicamentos.obter(med_id).quantidade_atual == 45             # nada mudou
    assert len(ctx.estoque.listar()) == antes                                # nada registrado


def test_passos_8_9_pesquisa_por_principio_ativo_e_por_localizacao(ctx):
    a = ctx.medicamentos.criar(paracetamol())                                # Corredor A
    b = ctx.medicamentos.criar(ibuprofeno())                                 # Corredor B
    c = ctx.medicamentos.criar(paracetamol(nome="Febrolyn", laboratorio="Outro Lab",
                                           lote="L002", corredor="A", estante="05", prateleira="01"))

    ids = {m.id for m in ctx.medicamentos.listar("paracetamol")}             # 8. princípio ativo
    assert ids == {a, c}
    assert {m.id for m in ctx.medicamentos.listar("PARACETAMOL")} == {a, c}
    assert {m.id for m in ctx.medicamentos.listar("analgesico")} == {a, c}   # sem acento
    assert {m.id for m in ctx.medicamentos.listar("paracetamol 750")} == {a, c}

    assert {m.id for m in ctx.medicamentos.listar("Corredor A")} == {a, c}   # 9. localização
    assert {m.id for m in ctx.medicamentos.listar("corredor b")} == {b}
    assert {m.id for m in ctx.medicamentos.listar("Estante 05")} == {c}
    assert {m.id for m in ctx.medicamentos.listar("prateleira 03")} == {a}
    assert ctx.medicamentos.listar("corredor z") == []


def test_passo_10_editar_medicamento_sem_mexer_no_estoque(ctx):
    med_id = ctx.medicamentos.criar(paracetamol(quantidade_atual=20))
    ctx.estoque.entrada(med_id, 5)
    med = ctx.medicamentos.obter(med_id)
    dados = paracetamol(nome="Paracetamol 750 mg Novo", corredor="C", estoque_minimo=7,
                        preco_venda_centavos=1000, quantidade_atual=9999)     # quantidade é ignorada
    ctx.medicamentos.atualizar(med_id, dados)
    novo = ctx.medicamentos.obter(med_id)
    assert novo.nome == "Paracetamol 750 mg Novo"
    assert novo.corredor == "C" and novo.estoque_minimo == 7
    assert novo.preco_venda_centavos == 1000
    assert novo.quantidade_atual == med.quantidade_atual == 25               # estoque preservado


def test_passos_11_12_excluir_mantem_historico(ctx):
    med_id = ctx.medicamentos.criar(paracetamol(quantidade_atual=20))
    outro = ctx.medicamentos.criar(ibuprofeno())
    ctx.estoque.entrada(med_id, 30, "Compra")
    ctx.estoque.saida(med_id, 5, "Venda")
    assert ctx.medicamentos.contar_movimentacoes(med_id) == 3                # inicial + 2

    ctx.medicamentos.excluir(med_id)                                         # 11. excluir
    assert ctx.medicamentos.obter(med_id) is None
    assert [m.id for m in ctx.medicamentos.listar()] == [outro]
    assert ctx.medicamentos.listar("paracetamol") == []
    with pytest.raises(NaoEncontradoError):
        ctx.estoque.entrada(med_id, 1)                                       # não movimenta excluído

    historico = [m for m in ctx.estoque.listar() if m.medicamento_id == med_id]   # 12. histórico
    assert len(historico) == 3
    assert all(m.excluido for m in historico)
    assert historico[0].medicamento == "Paracetamol 750 mg"


def test_passo_12_historico_tem_todos_os_campos_e_filtros(ctx):
    med_id = ctx.medicamentos.criar(paracetamol(quantidade_atual=20))
    outro = ctx.medicamentos.criar(ibuprofeno())
    ctx.estoque.entrada(med_id, 30, "Compra", "NF 123")
    ctx.estoque.saida(med_id, 5, "Venda")
    ctx.estoque.saida(outro, 2, "Perda ou avaria")

    movs = ctx.estoque.listar(medicamento_id=med_id)
    assert [(m.tipo, m.quantidade, m.estoque_anterior, m.estoque_posterior, m.motivo)
            for m in movs] == [("SAIDA", 5, 50, 45, "Venda"),
                               ("ENTRADA", 30, 20, 50, "Compra"),
                               ("ENTRADA", 20, 0, 20, "Estoque inicial")]
    assert movs[1].observacao == "NF 123"
    assert movs[0].data_hora.date() == HOJE and movs[0].quantidade_com_sinal == -5

    assert {m.tipo for m in ctx.estoque.listar(tipo="SAIDA")} == {"SAIDA"}
    assert len(ctx.estoque.listar(tipo="ENTRADA")) == 3      # 2 estoques iniciais + 1 compra
    assert len(ctx.estoque.listar(texto="ibuprofeno")) == 2
    assert len(ctx.estoque.listar(texto="ibuprofeno", tipo="SAIDA")) == 1
    assert len(ctx.estoque.listar(data_inicio=HOJE, data_fim=HOJE)) == 5
    assert ctx.estoque.listar(data_inicio=HOJE + timedelta(days=1)) == []
    assert ctx.estoque.listar(data_fim=HOJE - timedelta(days=1)) == []
    assert len(ctx.estoque.ultimas(3)) == 3


def test_passo_13_alertas_de_estoque_baixo(ctx):
    ok = ctx.medicamentos.criar(paracetamol(quantidade_atual=50, estoque_minimo=10))
    igual = ctx.medicamentos.criar(ibuprofeno(quantidade_atual=15, estoque_minimo=15))   # <= mínimo
    zerado = ctx.medicamentos.criar(ibuprofeno(nome="Ibuprofeno 400", lote="X", quantidade_atual=0))

    baixos = {m.id for m in ctx.medicamentos.listar(situacao=SIT_BAIXO)}
    assert baixos == {igual, zerado}
    assert ctx.relatorios.resumo().estoque_baixo == 2

    dias = ctx.config.dias_alerta
    assert ctx.medicamentos.obter(ok).alertas(HOJE, dias) == []
    assert ctx.medicamentos.obter(igual).alertas(HOJE, dias) == [ALERTA_ESTOQUE_BAIXO]
    assert ctx.medicamentos.obter(zerado).alertas(HOJE, dias) == [ALERTA_SEM_ESTOQUE]

    ctx.estoque.saida(ok, 41, "Venda")                        # 50 -> 9, passa a ser <= 10
    assert ok in {m.id for m in ctx.medicamentos.listar(situacao=SIT_BAIXO)}
    ctx.estoque.entrada(zerado, 100, "Compra")
    assert zerado not in {m.id for m in ctx.medicamentos.listar(situacao=SIT_BAIXO)}


def test_passo_14_alertas_de_validade_e_prazo_configuravel(ctx):
    vencido = ctx.medicamentos.criar(paracetamol(lote="V", data_validade=HOJE - timedelta(days=1)))
    hoje_v = ctx.medicamentos.criar(paracetamol(lote="H", data_validade=HOJE))
    em_20 = ctx.medicamentos.criar(paracetamol(lote="20", data_validade=HOJE + timedelta(days=20)))
    em_45 = ctx.medicamentos.criar(paracetamol(lote="45", data_validade=HOJE + timedelta(days=45)))
    longe = ctx.medicamentos.criar(paracetamol(lote="L", data_validade=HOJE + timedelta(days=400)))
    sem_data = ctx.medicamentos.criar(paracetamol(lote="S", data_validade=None))

    assert ctx.config.dias_alerta == 30                                       # padrão
    lm = ctx.medicamentos
    assert {m.id for m in lm.listar(situacao=SIT_VENCIDO)} == {vencido}
    assert {m.id for m in lm.listar(situacao=SIT_PROXIMO)} == {hoje_v, em_20}
    resumo = ctx.relatorios.resumo()
    assert (resumo.vencidos, resumo.proximos_vencimento) == (1, 2)

    ctx.config.definir_dias_alerta("60")                                      # configurável
    assert {m.id for m in lm.listar(situacao=SIT_PROXIMO, dias_alerta=60)} == {hoje_v, em_20, em_45}
    assert ctx.relatorios.resumo().proximos_vencimento == 3
    ctx.config.definir_dias_alerta("7")
    assert ctx.relatorios.resumo().proximos_vencimento == 1

    assert lm.obter(vencido).alertas(HOJE, 30) == [ALERTA_VENCIDO]
    assert lm.obter(em_20).alertas(HOJE, 30) == [ALERTA_PROXIMO_VENCIMENTO]
    assert lm.obter(longe).alertas(HOJE, 30) == []
    assert lm.obter(sem_data).alertas(HOJE, 30) == []
    assert [m.id for m in ctx.relatorios.lista_vencimentos(hoje=HOJE)][0] == vencido

    ctx.estoque.saida(vencido, 20, "Descarte por vencimento")                 # sem unidades: sem alerta
    assert ctx.relatorios.resumo().vencidos == 0


def test_passo_15_backup_e_restauracao(ctx, tmp_path):
    med_id = ctx.medicamentos.criar(paracetamol(quantidade_atual=20))
    ctx.estoque.entrada(med_id, 30)
    arquivo = tmp_path / "meu_backup.db"
    ctx.backup.criar_backup(arquivo)
    assert arquivo.exists() and arquivo.stat().st_size > 0

    # muda o sistema depois do backup
    ctx.estoque.saida(med_id, 10)
    novo = ctx.medicamentos.criar(ibuprofeno())
    ctx.medicamentos.excluir(med_id)
    assert ctx.medicamentos.obter(med_id) is None

    copia = ctx.backup.restaurar(arquivo)                                    # restaura
    assert copia.exists()                                                    # cópia de segurança
    med = ctx.medicamentos.obter(med_id)
    assert med is not None and med.quantidade_atual == 50                    # estado do backup
    assert ctx.medicamentos.obter(novo) is None                              # criado depois: sumiu
    assert len(ctx.estoque.listar(medicamento_id=med_id)) == 2

    # a cópia de segurança contém o estado anterior à restauração
    seguranca = sqlite3.connect(copia)
    total = seguranca.execute("SELECT COUNT(*) FROM medicamentos WHERE ativo = 1").fetchone()[0]
    seguranca.close()
    assert total == 1                                                        # só o ibuprofeno
    # o banco continua utilizável depois de restaurar
    ctx.estoque.entrada(med_id, 1)
    assert ctx.medicamentos.obter(med_id).quantidade_atual == 51


def test_restaurar_arquivos_invalidos(ctx, tmp_path):
    lixo = tmp_path / "lixo.db"
    lixo.write_text("isto não é um banco de dados")
    with pytest.raises(BackupError):
        ctx.backup.restaurar(lixo)
    outro = tmp_path / "outro.db"
    c = sqlite3.connect(outro)
    c.execute("CREATE TABLE coisa (x INTEGER)")
    c.commit()
    c.close()
    with pytest.raises(BackupError):
        ctx.backup.restaurar(outro)
    with pytest.raises(BackupError):
        ctx.backup.restaurar(tmp_path / "nao_existe.db")
    with pytest.raises(BackupError):
        ctx.backup.criar_backup(ctx.caminho_banco)
    # nada foi perdido com as tentativas
    assert ctx.medicamentos.total() == 0
    ctx.medicamentos.criar(paracetamol())
    assert ctx.medicamentos.total() == 1


# ================================================================ validações
def test_validacoes_de_cadastro(ctx):
    with pytest.raises(ValidacaoError) as e:
        ctx.medicamentos.criar(paracetamol(nome="   "))
    assert "nome" in e.value.erros
    with pytest.raises(ValidacaoError) as e:
        ctx.medicamentos.criar(paracetamol(principio_ativo=""))
    assert "principio_ativo" in e.value.erros
    with pytest.raises(ValidacaoError) as e:
        ctx.medicamentos.criar(paracetamol(quantidade_atual=-1, estoque_minimo=-3,
                                           preco_compra_centavos=-1))
    assert {"quantidade_atual", "estoque_minimo", "preco_compra"} <= set(e.value.erros)
    with pytest.raises(ValidacaoError) as e:
        ctx.medicamentos.criar(paracetamol(data_fabricacao=HOJE + timedelta(days=5)))
    assert "data_fabricacao" in e.value.erros
    with pytest.raises(ValidacaoError) as e:
        ctx.medicamentos.criar(paracetamol(data_fabricacao=date(2026, 5, 1),
                                           data_validade=date(2026, 4, 1)))
    assert "data_validade" in e.value.erros
    assert ctx.medicamentos.total() == 0                                     # nada foi gravado


def test_movimentacao_invalida(ctx):
    med_id = ctx.medicamentos.criar(paracetamol())
    for qtd in (0, -5):
        with pytest.raises(ValidacaoError):
            ctx.estoque.entrada(med_id, qtd)
    with pytest.raises(ValidacaoError):
        ctx.estoque.entrada(med_id, True)                                    # bool não vale
    with pytest.raises(ValidacaoError):
        ctx.estoque.entrada(med_id, 5, motivo="  ")
    with pytest.raises(ValidacaoError):
        ctx.estoque.registrar_movimentacao(med_id, "AJUSTE", 1, "x")
    with pytest.raises(NaoEncontradoError):
        ctx.estoque.saida(9999, 1)
    assert ctx.medicamentos.obter(med_id).quantidade_atual == 20


def test_estoque_negativo_impedido_pelo_proprio_banco(ctx):
    med_id = ctx.medicamentos.criar(paracetamol())
    with pytest.raises(sqlite3.IntegrityError):
        ctx.db.conn.execute("UPDATE medicamentos SET quantidade_atual = -1 WHERE id = ?", (med_id,))
    with pytest.raises(sqlite3.IntegrityError):                              # histórico incoerente
        ctx.db.conn.execute(
            "INSERT INTO movimentacoes (medicamento_id, medicamento_nome, tipo, quantidade, "
            "estoque_anterior, estoque_posterior, motivo, data_hora) "
            "VALUES (?, 'x', 'SAIDA', 5, 10, 99, 'x', '2026-01-01 10:00:00')", (med_id,))


def test_duplicidade_bloqueada_mas_lote_diferente_permitido(ctx):
    primeiro = ctx.medicamentos.criar(paracetamol())
    with pytest.raises(DuplicadoError) as e:
        ctx.medicamentos.criar(paracetamol(nome="  paracetamol   750 MG  ", laboratorio="lab teste"))
    assert e.value.existente_id == primeiro
    assert "Entrada" in str(e.value)
    ctx.medicamentos.criar(paracetamol(lote="L002"))                          # outro lote: ok
    ctx.medicamentos.criar(paracetamol(lote="L003", laboratorio="Outro"))
    assert ctx.medicamentos.total() == 3
    # editar para ficar igual a outro também é bloqueado
    outro = ctx.medicamentos.criar(paracetamol(lote="L004"))
    with pytest.raises(DuplicadoError):
        ctx.medicamentos.atualizar(outro, paracetamol(lote="L001"))
    ctx.medicamentos.atualizar(outro, paracetamol(lote="L004", estoque_minimo=3))   # a si mesmo: ok
    # depois de excluir, pode cadastrar de novo
    ctx.medicamentos.excluir(primeiro)
    ctx.medicamentos.criar(paracetamol())


def test_falha_no_meio_nao_deixa_registros_pela_metade(ctx):
    with pytest.raises(sqlite3.Error):
        with ctx.db.transacao():
            ctx.db.conn.execute("INSERT INTO categorias (nome) VALUES ('Teste')")
            ctx.db.conn.execute("INSERT INTO categorias (nome) VALUES ('teste')")   # duplicada
    assert ctx.medicamentos.listar_categorias() == []


def test_medicamentos_relacionados_sao_simetricos_e_baseados_nos_dados(ctx):
    a = ctx.medicamentos.criar(paracetamol())
    b = ctx.medicamentos.criar(paracetamol(nome="Febrolyn", lote="F1"))
    c = ctx.medicamentos.criar(ibuprofeno())
    assert ctx.medicamentos.obter(a).relacionados == []          # mesmo princípio ativo ≠ relação
    d = paracetamol(); d.relacionados = [b, a, b]                # a si mesmo e repetido: ignorados
    ctx.medicamentos.atualizar(a, d)
    assert [m.id for m in ctx.medicamentos.obter(a).relacionados] == [b]
    assert [m.id for m in ctx.medicamentos.obter(b).relacionados] == [a]      # via de mão dupla
    assert ctx.medicamentos.obter(c).relacionados == []
    ctx.medicamentos.excluir(b)
    assert ctx.medicamentos.obter(a).relacionados == []


def test_busca_por_codigo_lote_laboratorio_e_caracteres_especiais(ctx):
    a = ctx.medicamentos.criar(paracetamol(lote="ABC-123"))
    ctx.medicamentos.criar(ibuprofeno())
    assert [m.id for m in ctx.medicamentos.listar("abc-123")] == [a]
    assert [m.id for m in ctx.medicamentos.listar(f"{a:05d}")] == [a]
    assert a in [m.id for m in ctx.medicamentos.listar(str(a))]
    assert len(ctx.medicamentos.listar("lab teste")) == 2
    assert ctx.medicamentos.listar("100%") == [] and ctx.medicamentos.listar("_") == []
    assert [m.id for m in ctx.medicamentos.listar(categoria="analgesico")] == [a]
    assert len(ctx.medicamentos.listar(excluir_ids=(a,))) == 1


# =================================================================== conversores
@pytest.mark.parametrize("texto,esperado", [
    ("8,90", 890), ("8.90", 890), ("R$ 8,90", 890), ("1.234,56", 123456), ("1,234.56", 123456),
    ("1.234", 123400), ("10", 1000), ("0,5", 50), (",5", 50), ("", 0), ("  12  ", 1200)])
def test_converter_dinheiro(texto, esperado):
    assert converter_dinheiro(texto) == esperado


@pytest.mark.parametrize("texto", ["abc", "-5", "1,234,5", "8,905", "12,3,4", "R$ -1", "1e5"])
def test_converter_dinheiro_invalido(texto):
    with pytest.raises(ValueError):
        converter_dinheiro(texto)


def test_converter_inteiro():
    assert converter_inteiro("25") == 25 and converter_inteiro("") == 0
    for ruim in ("-1", "2,5", "abc", "1.000"):
        with pytest.raises(ValueError):
            converter_inteiro(ruim)
    with pytest.raises(ValueError):
        converter_inteiro("", obrigatorio=True)
    with pytest.raises(ValueError):
        converter_inteiro("0", obrigatorio=True, minimo=1)


def test_converter_data():
    assert converter_data("31/08/2027") == date(2027, 8, 31)
    assert converter_data("08/2027", fim_do_mes=True) == date(2027, 8, 31)
    assert converter_data("02/2028", fim_do_mes=True) == date(2028, 2, 29)     # bissexto
    assert converter_data("08/2027") == date(2027, 8, 1)
    assert converter_data("2027-08-31") == date(2027, 8, 31)
    assert converter_data("") is None
    for ruim in ("31/02/2027", "13/2027", "00/2027", "ontem", "1/1/27", "01/01/1800", "32/01/2027"):
        with pytest.raises(ValueError):
            converter_data(ruim)


# ================================================================== relatórios
def test_relatorios_e_exportacoes(ctx, tmp_path):
    a = ctx.medicamentos.criar(paracetamol(quantidade_atual=5, estoque_minimo=10,
                                           data_validade=HOJE - timedelta(days=3)))
    ctx.medicamentos.criar(ibuprofeno(data_validade=HOJE + timedelta(days=10)))
    ctx.estoque.entrada(a, 3, "Compra", "observação com acentuação: ção → ok")

    for tipo in (REL_ESTOQUE_ATUAL, REL_ESTOQUE_BAIXO, REL_VENCIDOS, REL_PROXIMOS,
                 REL_MOVIMENTACOES, REL_VALOR):
        rel = ctx.relatorios.gerar(tipo)
        assert rel.titulo and rel.colunas
        for linha in rel.linhas:
            assert len(linha) == len(rel.colunas)
        for nome, funcao, ext in (("csv", exportacao.exportar_csv, "csv"),
                                  ("xlsx", exportacao.exportar_excel, "xlsx"),
                                  ("pdf", exportacao.exportar_pdf, "pdf")):
            destino = tmp_path / f"{tipo}.{ext}"
            funcao(rel, destino, "Farmácia Teste")
            assert destino.stat().st_size > 100, (tipo, nome)

    valor = ctx.relatorios.gerar(REL_VALOR)
    assert dict(valor.resumo)["Valor total a preço de custo"] == "R$ 233,60"
    r = ctx.relatorios.resumo()
    assert r.total_medicamentos == 2 and r.total_unidades == 48
    assert r.valor_custo_centavos == 8 * 420 + 40 * 500
    assert r.valor_venda_centavos == 8 * 890 + 40 * 1100

    assert len(ctx.relatorios.gerar(REL_VENCIDOS).linhas) == 1
    assert len(ctx.relatorios.gerar(REL_PROXIMOS).linhas) == 1
    assert len(ctx.relatorios.gerar(REL_ESTOQUE_BAIXO).linhas) == 1

    # conteúdo do CSV e do Excel
    csv_txt = (tmp_path / f"{REL_ESTOQUE_ATUAL}.csv").read_text(encoding="utf-8-sig")
    assert csv_txt.splitlines()[0].startswith("Código;Medicamento;")
    assert "Paracetamol 750 mg" in csv_txt and "4,20" not in csv_txt.split("\n")[0]
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / f"{REL_VALOR}.xlsx").active
    assert ws["A2"].value == "Valor estimado do estoque"
    valores = [c.value for c in ws[6]]                     # primeira linha de dados (após o cabeçalho)
    assert valores[2] in (8, 40)                           # estoque numérico, não texto


def test_relatorio_movimentacoes_respeita_periodo(ctx):
    a = ctx.medicamentos.criar(paracetamol())
    ctx.estoque.registrar_movimentacao(a, "SAIDA", 3, "Venda",
                                       data_hora=datetime.now() - timedelta(days=40))
    rel = ctx.relatorios.gerar(REL_MOVIMENTACOES, data_inicio=HOJE - timedelta(days=7), data_fim=HOJE)
    assert len(rel.linhas) == 1                                              # só o estoque inicial
    assert len(ctx.relatorios.gerar(REL_MOVIMENTACOES).linhas) == 2


# ================================================================ demonstração
def test_dados_de_demonstracao(ctx):
    assert not ctx.demo.existem_dados_demo()
    real = ctx.medicamentos.criar(paracetamol(nome="Produto real", lote="R1"))
    qtd = ctx.demo.adicionar()
    assert qtd == 15 and ctx.demo.existem_dados_demo()
    meds = ctx.medicamentos.listar()
    assert len(meds) == 16
    assert all("FICTÍCIO" in m.observacoes for m in meds if m.demo)
    resumo = ctx.relatorios.resumo()
    assert resumo.estoque_baixo >= 3 and resumo.vencidos >= 1 and resumo.proximos_vencimento >= 2
    gen = next(m for m in meds if m.nome == "Paracetamol Genérico" and m.dosagem == "750")
    nomes = {r.nome for r in ctx.medicamentos.obter(gen.id).relacionados}
    assert nomes == {"Febrolyn", "Paracetamol Genérico"}          # marca + gotas, definidos na demo
    assert len(ctx.estoque.listar()) > 30                                    # histórico gerado
    # histórico coerente: cada movimento parte do estoque anterior do movimento anterior
    for m in meds:
        if m.demo:
            movs = list(reversed(ctx.estoque.listar(medicamento_id=m.id)))
            assert movs[-1].estoque_posterior == m.quantidade_atual
            for anterior, atual in zip(movs, movs[1:]):
                assert anterior.estoque_posterior == atual.estoque_anterior

    assert ctx.demo.remover() == 15
    assert not ctx.demo.existem_dados_demo()
    restantes = ctx.medicamentos.listar()
    assert [m.id for m in restantes] == [real]                               # o real fica intacto
    assert ctx.medicamentos.obter(real).quantidade_atual == 20
    assert len(ctx.estoque.listar()) == 1
    assert ctx.medicamentos.listar_categorias() == ["Analgésico"]


def test_configuracoes(ctx):
    assert ctx.config.dias_alerta == 30 and ctx.config.nome_farmacia == ""
    for ruim in ("0", "-3", "abc", "", "9999", "2,5"):
        with pytest.raises(ValidacaoError):
            ctx.config.definir_dias_alerta(ruim)
    assert ctx.config.definir_dias_alerta(" 45 ") == 45 and ctx.config.dias_alerta == 45
    assert ctx.config.definir_nome_farmacia("  Farmácia   Saúde  ") == "Farmácia Saúde"
    with pytest.raises(ValidacaoError):
        ctx.config.definir_nome_farmacia("x" * 61)
